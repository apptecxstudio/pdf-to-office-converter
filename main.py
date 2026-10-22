import os
import shutil
import uuid
from fastapi import FastAPI, UploadFile, File, Request, HTTPException
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pdf2docx import Converter
import pdfplumber
import openpyxl
from pptx import Presentation
from pptx.util import Inches
import fitz  # PyMuPDF

app = FastAPI(title="PDF to Office Converter API")

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)
app.mount("/uploads", StaticFiles(directory=UPLOAD_DIR), name="uploads")

@app.get("/")
def home():
    return {"status": "ok", "message": "PDF Converter API is active and ready"}

# 1. API for uploading file (Used in your Flutter app)
@app.post("/upload_files")
async def upload_files(request: Request):
    form = await request.form()
    # Flutter sends multi-files or file
    uploaded_file = form.get("multi-files") or form.get("file")
    if not uploaded_file:
        raise HTTPException(status_code=400, detail="No file uploaded")
    
    clean_filename = f"{uuid.uuid4().hex[:8]}_{uploaded_file.filename}"
    file_path = os.path.join(UPLOAD_DIR, clean_filename)
    
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(uploaded_file.file, buffer)
        
    return {
        "status": True,
        "files": [
            {
                "path": file_path,
                "filename": clean_filename
            }
        ]
    }

# 2. PDF to Word (DOCX) Endpoint
@app.post("/pdftodoc/docx")
async def convert_pdf_to_docx(payload: dict):
    pdf_rel_path = payload.get("path")
    if not pdf_rel_path or not os.path.exists(pdf_rel_path):
        raise HTTPException(status_code=404, detail="PDF file not found")
        
    docx_path = os.path.splitext(pdf_rel_path)[0] + ".docx"
    try:
        cv = Converter(pdf_rel_path)
        cv.convert(docx_path, start=0, end=None)
        cv.close()
        return {"status": True, "path": docx_path}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# 3. PDF to Excel Endpoint
@app.post("/convert-pdf-to-excel/")
async def convert_pdf_to_excel(request: Request, file: UploadFile = File(...)):
    clean_filename = f"{uuid.uuid4().hex[:8]}_{file.filename}"
    pdf_path = os.path.join(UPLOAD_DIR, clean_filename)
    excel_path = os.path.splitext(pdf_path)[0] + ".xlsx"
    
    with open(pdf_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            clean_row = [cell if cell is not None else "" for cell in row]
                            ws.append(clean_row)
                else:
                    text = page.extract_text()
                    if text:
                        for line in text.split("\n"):
                            ws.append([line])
        wb.save(excel_path)
        base_url = str(request.base_url).rstrip("/")
        download_url = f"{base_url}/{excel_path}"
        return {"download_link": download_url}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

# 4. PDF to PPTX Endpoint
@app.post("/convert_pdf_to_pptx/")
async def convert_pdf_to_pptx(file: UploadFile = File(...)):
    clean_filename = f"{uuid.uuid4().hex[:8]}_{file.filename}"
    pdf_path = os.path.join(UPLOAD_DIR, clean_filename)
    pptx_path = os.path.splitext(pdf_path)[0] + ".pptx"
    
    with open(pdf_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    try:
        prs = Presentation()
        # Set slide size to 16:9 widescreen
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_slide_layout = prs.slide_layouts[6]
        
        doc = fitz.open(pdf_path)
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            pix = page.get_pixmap(dpi=150)
            img_path = f"{pdf_path}_page_{page_num}.png"
            pix.save(img_path)
            
            slide = prs.slides.add_slide(blank_slide_layout)
            slide.shapes.add_picture(img_path, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
            if os.path.exists(img_path):
                os.remove(img_path)
        doc.close()
        prs.save(pptx_path)
        return FileResponse(pptx_path, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", filename=os.path.basename(pptx_path))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
